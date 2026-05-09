import csv
import io

from django.http import HttpResponse
from django.db.models import Q
from django.utils import timezone
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from rest_framework import status
from rest_framework.authentication import TokenAuthentication
from rest_framework.decorators import api_view, authentication_classes, permission_classes
from rest_framework.permissions import IsAdminUser, IsAuthenticated
from rest_framework.response import Response
from applications.globals.services import get_extra_info_by_user
from applications.complaint_system.models import ComplaintStatus
from . import messages, selectors, services
from .permissions import IsCaretaker, IsSupervisor, IsSupervisorOrAdmin, IsComplaintAdminOrServiceAuthority
from .permissions import (
    IsCaretaker,
    IsSupervisor,
    IsSupervisorOrAdmin,
    IsComplaintAdminOrServiceAuthority,
    IsSupervisorOrComplaintAdmin,
)
from .serializers import (
    CaretakerSerializer,
    ComplaintAdminAssignSerializer,
    ComplaintCloseSerializer,
    ComplaintCreateSerializer,
    ComplaintEscalateSerializer,
    ComplaintFeedbackSerializer,
    ComplaintProgressSerializer,
    ComplaintReopenSerializer,
    ReopenRequestCreateSerializer,
    ReopenRequestReviewSerializer,
    ReopenRequestSerializer,
    ReportExportSerializer,
    SupervisorReassignSerializer,
    StudentComplainSerializer,
    SupervisorSerializer,
    WorkerSerializer,
)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
@authentication_classes([TokenAuthentication])
def complaint_details_api(request, detailcomp_id1):
    complaint_detail = selectors.get_complaint(detailcomp_id1)
    if not services.can_access_complaint(request.user, complaint_detail):
        return Response({'message': messages.PERMISSION_DENIED}, status=status.HTTP_403_FORBIDDEN)
    response = services.build_complaint_payload(complaint_detail)
    return Response(data=response, status=status.HTTP_200_OK)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
@authentication_classes([TokenAuthentication])
def student_complain_api(request):
    extra_info = get_extra_info_by_user(request.user)
    complain = services.get_complains_for_user(extra_info)
    complains = StudentComplainSerializer(complain, many=True).data
    resp = {
        'student_complain': complains,
    }
    return Response(data=resp, status=status.HTTP_200_OK)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
@authentication_classes([TokenAuthentication])
def create_complain_api(request):
    serializer = ComplaintCreateSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    extra_info = get_extra_info_by_user(request.user)
    idempotency_key = request.META.get('HTTP_X_IDEMPOTENCY_KEY', '').strip() or None
    data = services.create_complaint(extra_info, serializer.validated_data, idempotency_key=idempotency_key)
    return Response(data, status=status.HTTP_201_CREATED)


@api_view(['DELETE', 'PUT'])
@permission_classes([IsAuthenticated])
@authentication_classes([TokenAuthentication])
def edit_complain_api(request, c_id):
    complain = selectors.get_complaint(c_id)
    if not services.can_access_complaint(request.user, complain):
        return Response({'message': messages.PERMISSION_DENIED}, status=status.HTTP_403_FORBIDDEN)
    if request.method == 'DELETE':
        services.delete_complaint(complain)
        return Response(status=status.HTTP_204_NO_CONTENT)
    elif request.method == 'PUT':
        data = services.update_complaint(complain, request.data)
        return Response(data, status=status.HTTP_200_OK)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
@authentication_classes([TokenAuthentication])
def update_progress_api(request, c_id):
    complain = selectors.get_complaint(c_id)
    if not services.can_access_complaint(request.user, complain):
        return Response({'message': messages.PERMISSION_DENIED}, status=status.HTTP_403_FORBIDDEN)
    serializer = ComplaintProgressSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    services.update_progress(
        complain,
        get_extra_info_by_user(request.user),
        serializer.validated_data['status'],
        serializer.validated_data.get('note', ''),
        request.FILES.get('upload_resolved'),
    )
    return Response({'message': 'Progress updated.'}, status=status.HTTP_200_OK)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
@authentication_classes([TokenAuthentication])
def escalate_complaint_api(request, c_id):
    complain = selectors.get_complaint(c_id)
    serializer = ComplaintEscalateSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    services.escalate_complaint(
        complain,
        get_extra_info_by_user(request.user),
        serializer.validated_data['justification'],
        is_auto=False,
    )
    return Response({'message': 'Complaint escalated.'}, status=status.HTTP_200_OK)


@api_view(['POST'])
@permission_classes([IsAuthenticated, IsSupervisor])
@authentication_classes([TokenAuthentication])
def supervisor_reassign_api(request, c_id):
    complain = selectors.get_complaint(c_id)
    if not services.can_access_complaint(request.user, complain):
        return Response({'message': messages.PERMISSION_DENIED}, status=status.HTTP_403_FORBIDDEN)

    serializer = SupervisorReassignSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)

    caretaker = selectors.get_caretaker(serializer.validated_data['caretaker_id'])
    services.supervisor_reassign(
        complain,
        actor=get_extra_info_by_user(request.user),
        caretaker=caretaker,
        note=serializer.validated_data.get('note', ''),
    )
    return Response({'message': 'Complaint reassigned by supervisor.'}, status=status.HTTP_200_OK)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
@authentication_classes([TokenAuthentication])
def close_complaint_api(request, c_id):
    complain = selectors.get_complaint(c_id)
    if not services.can_access_complaint(request.user, complain):
        return Response({'message': messages.PERMISSION_DENIED}, status=status.HTTP_403_FORBIDDEN)
    serializer = ComplaintCloseSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    services.close_complaint(complain, get_extra_info_by_user(request.user), serializer.validated_data['verified'])
    return Response({'message': 'Complaint closed.'}, status=status.HTTP_200_OK)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
@authentication_classes([TokenAuthentication])
def reopen_complaint_api(request, c_id):
    """BR-CM-008 fix: Direct reopen now creates a supervised reopen request
    instead of bypassing supervisor approval."""
    complain = selectors.get_complaint(c_id)
    if not services.can_access_complaint(request.user, complain):
        return Response({'message': messages.PERMISSION_DENIED}, status=status.HTTP_403_FORBIDDEN)
    serializer = ComplaintReopenSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    extra_info = get_extra_info_by_user(request.user)
    if complain.complainer_id != extra_info.id and not (request.user.is_superuser or services._is_complaint_admin(extra_info) or services._is_service_authority(extra_info)):
        return Response(
            {'message': 'Only the original complainant can request a reopen.'},
            status=status.HTTP_403_FORBIDDEN,
        )
    reopen_req = services.create_reopen_request(
        complain, extra_info, serializer.validated_data['justification']
    )
    return Response(
        {
            'message': 'Reopen request submitted for supervisor approval.',
            'request_id': reopen_req.id,
        },
        status=status.HTTP_201_CREATED,
    )


@api_view(['GET', 'POST'])
@permission_classes([IsAuthenticated])
@authentication_classes([TokenAuthentication])
def reopen_request_api(request, c_id):
    """UC-CM-011: Submit a reopen request (POST by complainant) or review it.
    WF-401 fix: Identity verification enforced — only complainant can create,
    only supervisor/admin can review."""
    complain = selectors.get_complaint(c_id)
    if request.method == 'GET':
        # List reopen requests for this complaint
        if not services.can_access_complaint(request.user, complain):
            return Response({'message': messages.PERMISSION_DENIED}, status=status.HTTP_403_FORBIDDEN)
        requests = complain.reopen_requests.order_by('-created_at')
        data = ReopenRequestSerializer(requests, many=True).data
        return Response({'reopen_requests': data}, status=status.HTTP_200_OK)

    # POST: either create a new request or review an existing one
    extra_info = get_extra_info_by_user(request.user)
    action = request.data.get('action', 'create')

    if action == 'create':
        # WF-401: Identity verification — only the original complainant can request reopen
        if complain.complainer_id != extra_info.id and not request.user.is_superuser:
            return Response(
                {'message': 'Only the original complainant can request a reopen.'},
                status=status.HTTP_403_FORBIDDEN,
            )
        serializer = ReopenRequestCreateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        reopen_req = services.create_reopen_request(complain, extra_info, serializer.validated_data['justification'])
        return Response(ReopenRequestSerializer(reopen_req).data, status=status.HTTP_201_CREATED)

    elif action == 'review':
        # WF-401: Identity verification — only supervisor or admin can review
        if not (services._is_supervisor(extra_info) or request.user.is_superuser):
            return Response(
                {'message': 'Only supervisors or admins can review reopen requests.'},
                status=status.HTTP_403_FORBIDDEN,
            )
        if services._is_supervisor(extra_info):
            from applications.complaint_system.models import Supervisor as SupervisorModel
            supervisor_obj = SupervisorModel.objects.filter(sup_id=extra_info).first()
            if not supervisor_obj or complain.assigned_supervisor_id != supervisor_obj.id:
                return Response(
                    {'message': 'Only the assigned supervisor can review this reopen request.'},
                    status=status.HTTP_403_FORBIDDEN,
                )
        request_id = request.data.get('request_id')
        if not request_id:
            return Response({'message': 'request_id is required.'}, status=status.HTTP_400_BAD_REQUEST)
        from applications.complaint_system.models import ReopenRequest as ReopenRequestModel
        try:
            reopen_req = ReopenRequestModel.objects.get(id=request_id, complaint=complain)
        except ReopenRequestModel.DoesNotExist:
            return Response({'message': 'Reopen request not found.'}, status=status.HTTP_404_NOT_FOUND)
        serializer = ReopenRequestReviewSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        services.approve_reopen_request(
            reopen_req, extra_info,
            serializer.validated_data['approved'],
            serializer.validated_data.get('review_note', ''),
        )
        return Response({'message': 'Reopen request reviewed.'}, status=status.HTTP_200_OK)

    return Response({'message': 'Invalid action.'}, status=status.HTTP_400_BAD_REQUEST)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
@authentication_classes([TokenAuthentication])
def feedback_api(request, c_id):
    complain = selectors.get_complaint(c_id)
    if not services.can_access_complaint(request.user, complain):
        return Response({'message': messages.PERMISSION_DENIED}, status=status.HTTP_403_FORBIDDEN)
    serializer = ComplaintFeedbackSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    services.submit_feedback(complain, serializer.validated_data['rating'], serializer.validated_data.get('comments', ''))
    return Response({'message': 'Feedback submitted.'}, status=status.HTTP_200_OK)


@api_view(['GET'])
@permission_classes([IsAuthenticated, IsSupervisorOrComplaintAdmin])
@authentication_classes([TokenAuthentication])
def report_api(request):
    """WF-601: Report API with date range validation."""
    from datetime import datetime
    status_val = request.GET.get('status')

    # Parse and validate date inputs
    start_date_str = request.GET.get('start_date') or None
    end_date_str = request.GET.get('end_date') or None
    start_date = None
    end_date = None
    if start_date_str:
        try:
            start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date()
        except ValueError:
            return Response(
                {'message': 'Invalid start_date format. Use YYYY-MM-DD.'},
                status=status.HTTP_400_BAD_REQUEST,
            )
    if end_date_str:
        try:
            end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date()
        except ValueError:
            return Response(
                {'message': 'Invalid end_date format. Use YYYY-MM-DD.'},
                status=status.HTTP_400_BAD_REQUEST,
            )
    if start_date and end_date and start_date > end_date:
        return Response(
            {'message': 'start_date must be on or before end_date.'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    filters = {
        'location': request.GET.get('location') or None,
        'complaint_type': request.GET.get('complaint_type') or None,
        'priority': request.GET.get('priority') or None,
        'status': None,
        'start_date': start_date,
        'end_date': end_date,
    }
    if status_val and status_val.strip():
        try:
            filters['status'] = int(status_val)
        except ValueError:
            return Response({'message': 'Invalid status filter.'}, status=status.HTTP_400_BAD_REQUEST)
    qs = services.generate_report(filters)
    if not request.user.is_superuser:
        allowed_ids = [c.id for c in qs if services.can_access_complaint(request.user, c)]
        qs = qs.filter(id__in=allowed_ids)
    data = StudentComplainSerializer(qs, many=True).data
    summary = services.build_report_summary(qs)
    return Response({'results': data, 'count': len(data), 'summary': summary}, status=status.HTTP_200_OK)


@api_view(['GET'])
@permission_classes([IsAuthenticated, IsSupervisorOrComplaintAdmin])
@authentication_classes([TokenAuthentication])
def export_report_api(request):
    """UC-CM-008: Export report as CSV/PDF/Excel."""
    # NOTE: `format` is reserved by DRF content negotiation. Using it as a
    # query param can trigger a 404 before this view executes.
    export_format = request.GET.get('export_format') or request.GET.get('file_format') or 'csv'
    status_val = request.GET.get('status')
    filters = {
        'location': request.GET.get('location') or None,
        'complaint_type': request.GET.get('complaint_type') or None,
        'priority': request.GET.get('priority') or None,
        'status': None,
        'start_date': request.GET.get('start_date') or None,
        'end_date': request.GET.get('end_date') or None,
    }
    if status_val and status_val.strip():
        try:
            filters['status'] = int(status_val)
        except ValueError:
            return Response({'message': 'Invalid status filter.'}, status=status.HTTP_400_BAD_REQUEST)

    qs = services.generate_report(filters)
    if not request.user.is_superuser:
        allowed_ids = [c.id for c in qs if services.can_access_complaint(request.user, c)]
        qs = qs.filter(id__in=allowed_ids)
    summary = services.build_report_summary(qs)

    if export_format == 'csv':
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="complaint_report.csv"'
        writer = csv.writer(response)
        writer.writerow(['ID', 'Type', 'Location', 'Priority', 'Status', 'Date', 'SLA Deadline', 'Details'])
        status_map = dict(services.ComplaintStatus.choices)
        for c in qs:
            writer.writerow([
                c.id, c.complaint_type, c.location, c.priority,
                status_map.get(c.status, c.status),
                c.complaint_date.strftime('%Y-%m-%d %H:%M') if c.complaint_date else 'N/A',
                c.sla_deadline.strftime('%Y-%m-%d %H:%M') if getattr(c, 'sla_deadline', None) else 'N/A',
                c.details,
            ])
        # Write summary
        writer.writerow([])
        writer.writerow(['--- KPI Summary ---'])
        writer.writerow(['Total Complaints', summary['total']])
        writer.writerow(['Avg Resolution (hours)', summary.get('avg_resolution_hours', 'N/A')])
        writer.writerow(['Reopen Rate (%)', summary.get('reopen_rate', 0)])
        writer.writerow(['SLA Compliance (%)', summary.get('sla_compliance', 100)])
        return response

    elif export_format == 'excel':
        # True .xlsx export using openpyxl
        wb = openpyxl.Workbook()
        ws_data = wb.active
        ws_data.title = 'Complaint Report'

        # Header styling
        header_font = Font(bold=True, color='FFFFFF', size=11)
        header_fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin'),
        )

        headers = ['ID', 'Type', 'Location', 'Priority', 'Status', 'Date', 'SLA Deadline', 'Details']
        for col_num, header in enumerate(headers, 1):
            cell = ws_data.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        status_map = dict(services.ComplaintStatus.choices)
        for row_num, c in enumerate(qs, 2):
            row_data = [
                c.id, c.complaint_type, c.location, c.priority,
                status_map.get(c.status, c.status),
                c.complaint_date.strftime('%Y-%m-%d %H:%M') if getattr(c, 'complaint_date', None) else 'N/A',
                c.sla_deadline.strftime('%Y-%m-%d %H:%M') if getattr(c, 'sla_deadline', None) else 'N/A',
                c.details,
            ]
            for col_num, value in enumerate(row_data, 1):
                cell = ws_data.cell(row=row_num, column=col_num, value=value)
                cell.border = thin_border

        # Auto-adjust column widths
        for col in ws_data.columns:
            max_length = max((len(str(cell.value or '')) for cell in col), default=10)
            ws_data.column_dimensions[col[0].column_letter].width = min(max_length + 4, 40)

        # KPI Summary sheet
        ws_summary = wb.create_sheet('KPI Summary')
        summary_header_fill = PatternFill(start_color='1ABC9C', end_color='1ABC9C', fill_type='solid')
        kpi_data = [
            ['KPI', 'Value'],
            ['Total Complaints', summary['total']],
            ['Avg Resolution (hours)', summary.get('avg_resolution_hours', 'N/A')],
            ['Reopen Rate (%)', summary.get('reopen_rate', 0)],
            ['SLA Compliance (%)', summary.get('sla_compliance', 100)],
        ]
        for row_num, row in enumerate(kpi_data, 1):
            for col_num, value in enumerate(row, 1):
                cell = ws_summary.cell(row=row_num, column=col_num, value=value)
                cell.border = thin_border
                if row_num == 1:
                    cell.font = header_font
                    cell.fill = summary_header_fill
                    cell.alignment = header_alignment
        ws_summary.column_dimensions['A'].width = 25
        ws_summary.column_dimensions['B'].width = 20

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = 'attachment; filename="complaint_report.xlsx"'
        return response

    elif export_format == 'pdf':
        # True PDF export using ReportLab
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), title='Complaint Management Report')
        styles = getSampleStyleSheet()
        elements = []

        # Title
        title_style = styles['Title']
        elements.append(Paragraph('Complaint Management Report', title_style))
        elements.append(Spacer(1, 0.3 * inch))

        # KPI Summary table
        elements.append(Paragraph('KPI Summary', styles['Heading2']))
        kpi_table_data = [
            ['Metric', 'Value'],
            ['Total Complaints', str(summary['total'])],
            ['Avg Resolution Time', f"{summary.get('avg_resolution_hours', 'N/A')} hours"],
            ['Reopen Rate', f"{summary.get('reopen_rate', 0)}%"],
            ['SLA Compliance', f"{summary.get('sla_compliance', 100)}%"],
        ]
        kpi_table = Table(kpi_table_data, colWidths=[2.5 * inch, 2 * inch])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2C3E50')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 0.4 * inch))

        # Complaint data table
        elements.append(Paragraph('Complaint Details', styles['Heading2']))
        status_map = dict(services.ComplaintStatus.choices)
        data_header = ['ID', 'Type', 'Location', 'Priority', 'Status', 'Date', 'Details']
        table_data = [data_header]
        for c in qs:
            table_data.append([
                str(c.id),
                str(c.complaint_type),
                str(c.location),
                str(c.priority),
                str(status_map.get(c.status, c.status)),
                c.complaint_date.strftime('%Y-%m-%d') if getattr(c, 'complaint_date', None) else 'N/A',
                str(c.details)[:60],
            ])

        col_widths = [0.5 * inch, 1 * inch, 1.2 * inch, 0.9 * inch, 0.9 * inch, 1 * inch, 3.5 * inch]
        complaint_table = Table(table_data, colWidths=col_widths, repeatRows=1)
        complaint_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2C3E50')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
        ]))
        elements.append(complaint_table)

        doc.build(elements)
        buffer.seek(0)
        response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="complaint_report.pdf"'
        return response

    return Response({'message': 'Unsupported format.'}, status=status.HTTP_400_BAD_REQUEST)


@api_view(['GET'])
@permission_classes([IsAuthenticated, IsSupervisor])
@authentication_classes([TokenAuthentication])
def supervisor_dashboard_api(request):
    """Supervisor queue data: assigned, escalated, reopen requested."""
    extra_info = get_extra_info_by_user(request.user)
    data = services.get_supervisor_dashboard(extra_info)
    return Response(data, status=status.HTTP_200_OK)


@api_view(['GET', 'POST'])
@permission_classes([IsAuthenticated, IsComplaintAdminOrServiceAuthority])
@authentication_classes([TokenAuthentication])
def admin_oversight_api(request):
    if request.method == 'GET':
        qs = selectors.list_unresolved()
        scope = (request.GET.get('scope') or 'overdue_escalated').strip().lower()
        if scope != 'all_unresolved':
            now = timezone.now()
            qs = qs.filter(
                Q(status=ComplaintStatus.ESCALATED)
                | Q(status__in=[ComplaintStatus.PENDING, ComplaintStatus.IN_PROGRESS, ComplaintStatus.REOPENED], sla_deadline__lt=now)
            )
        data = StudentComplainSerializer(qs, many=True).data
        return Response({'results': data, 'count': len(data)}, status=status.HTTP_200_OK)

    serializer = ComplaintAdminAssignSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    complaint_id = serializer.validated_data['complaint_id']
    complaint = selectors.get_complaint(complaint_id)
    caretaker = None
    supervisor = None
    if serializer.validated_data.get('caretaker_id') is not None:
        caretaker = selectors.get_caretaker(serializer.validated_data['caretaker_id'])
    if serializer.validated_data.get('supervisor_id') is not None:
        supervisor = selectors.get_supervisor(serializer.validated_data['supervisor_id'])
    services.admin_assign(complaint, caretaker=caretaker, supervisor=supervisor, actor=get_extra_info_by_user(request.user))
    return Response({'message': 'Admin assignment updated.'}, status=status.HTTP_200_OK)


@api_view(['GET', 'POST'])
@permission_classes([IsAuthenticated, IsCaretaker])
@authentication_classes([TokenAuthentication])
def worker_api(request):
    if request.method == 'GET':
        worker = services.list_workers()
        workers = WorkerSerializer(worker, many=True).data
        resp = {
            'workers': workers,
        }
        return Response(data=resp, status=status.HTTP_200_OK)

    elif request.method == 'POST':
        data = services.create_worker(request.data)
        return Response(data, status=status.HTTP_201_CREATED)


@api_view(['DELETE', 'PUT'])
@permission_classes([IsAuthenticated, IsCaretaker])
@authentication_classes([TokenAuthentication])
def edit_worker_api(request, w_id):
    worker = selectors.get_worker(w_id)
    if request.method == 'DELETE':
        services.delete_worker(worker)
        return Response(status=status.HTTP_204_NO_CONTENT)
    elif request.method == 'PUT':
        data = services.update_worker(worker, request.data)
        return Response(data, status=status.HTTP_200_OK)


@api_view(['GET', 'POST'])
@permission_classes([IsAuthenticated, IsSupervisorOrComplaintAdmin])
@authentication_classes([TokenAuthentication])
def caretaker_api(request):
    if request.method == 'GET':
        caretaker = services.list_caretakers()
        caretakers = CaretakerSerializer(caretaker, many=True).data
        resp = {
            'caretakers': caretakers,
        }
        return Response(data=resp, status=status.HTTP_200_OK)

    elif request.method == 'POST':
        data = services.create_caretaker(request.data)
        return Response(data, status=status.HTTP_201_CREATED)


@api_view(['DELETE', 'PUT'])
@permission_classes([IsAuthenticated, IsSupervisorOrComplaintAdmin])
@authentication_classes([TokenAuthentication])
def edit_caretaker_api(request, c_id):
    caretaker = selectors.get_caretaker(c_id)
    if request.method == 'DELETE':
        services.delete_caretaker(caretaker)
        return Response(status=status.HTTP_204_NO_CONTENT)
    elif request.method == 'PUT':
        data = services.update_caretaker(caretaker, request.data)
        return Response(data, status=status.HTTP_200_OK)


@api_view(['GET', 'POST'])
@permission_classes([IsAuthenticated, IsSupervisorOrComplaintAdmin])
@authentication_classes([TokenAuthentication])
def supervisor_api(request):
    if request.method == 'GET':
        supervisor = services.list_supervisors()
        supervisors = SupervisorSerializer(supervisor, many=True).data
        resp = {
            'supervisors': supervisors,
        }
        return Response(data=resp, status=status.HTTP_200_OK)

    elif request.method == 'POST':
        data = services.create_supervisor(request.data)
        return Response(data, status=status.HTTP_201_CREATED)


@api_view(['DELETE', 'PUT'])
@permission_classes([IsAuthenticated, IsSupervisorOrComplaintAdmin])
@authentication_classes([TokenAuthentication])
def edit_supervisor_api(request, s_id):
    supervisor = selectors.get_supervisor(s_id)
    if request.method == 'DELETE':
        services.delete_supervisor(supervisor)
        return Response(status=status.HTTP_204_NO_CONTENT)
    elif request.method == 'PUT':
        data = services.update_supervisor(supervisor, request.data)
        return Response(data, status=status.HTTP_200_OK)
